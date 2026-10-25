"""Builds TechVestVentures American Renaissance Graphs from graphs/American Industrial Renaissance (PNGs + pptx speaker notes). Text is verbatim from the deck."""
import os,re,sys,glob,shutil,hashlib
from pptx import Presentation
sys.path.insert(0,'/workspace/risk-watcher/graphs-site');os.chdir(os.path.dirname(os.path.abspath(__file__)))
import common
from common import *
T="TechVestVentures American Renaissance Graphs";F="/workspace/risk-watcher/graphs/American Industrial Renaissance"
png={f.split("/")[-1][:2]:f for f in glob.glob(F+"/*.png")}
H={hashlib.md5(open(f,'rb').read()).hexdigest():k for k,f in png.items()}
G={};gaps=[]
P=Presentation(F+"/American Industrial Renaissance.pptx")
URL=re.compile(r'https?://\S+')
for s in P.slides:
    pics=[sh for sh in s.shapes if sh.shape_type==13]
    if not pics: continue
    n=s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''
    txt=[sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    m=re.search(r'^File: (\d\d) ',n,re.M)
    k=H.get(hashlib.md5(pics[0].image.blob).hexdigest()) or (m[1] if m else None)
    if not k: gaps.append("unmapped slide: "+n[:60]); continue
    d={"title":(re.search(r'^Title: (.*)',n,re.M) or [None,None])[1],"caption":(re.search(r'^Caption: (.*)',n,re.M) or [None,txt[0] if txt else ''])[1],
       "label":(re.search(r'^Label: (.*)',n,re.M) or [None,''])[1],"src":re.findall(r'^Source: (\S+)',n,re.M),"bul":[]}
    if not d["title"]: d["title"]=png[k].split("/")[-1][3:-4]
    for l in n.split("\n"):
        l=l.strip()
        if l.startswith("•") or l.startswith("Verdict:"):
            us=URL.findall(l);t=l.lstrip("• ").strip()
            d["bul"].append((t,us))
            if not us and l.startswith("•"): gaps.append(f"PNG {k}: bullet without link: {t}")
    G[k]=d
TABS=[("1","Total Investment Pledges","10 11 12 26 02 04 07 08"),("2","Semiconductors","16 05"),("3","Factory Construction and Capacity","15 18 03 06"),
("4","Manufacturing Output and Recessions","17 28 29 22"),("5","Jobs and Reshoring","09 19 24"),("6","Foreign Investment Flows","13 14 27"),("7","Clean-Tech Manufacturing","25"),("8","Global Comparison","20 21 23"),("9","Pharmaceuticals","30 31 32"),("10","Autos and EVs","33 34 35"),("11","Steel and Metals","36 37 38")]
common.CSS+="""
figure.g{display:grid;grid-template-columns:minmax(0,7fr) minmax(260px,3fr);gap:22px;align-items:start;margin:0 0 40px}
figure.g img{width:100%;height:auto;max-width:none}
figure.g figcaption{font-size:18px;line-height:1.5;margin:0}figure.g figcaption p{margin:0 0 12px}
figure.g figcaption .t{font-weight:700;display:block;margin-bottom:8px}figure.g ul{margin:0 0 12px;padding-left:1.1em;font-size:16px}figure.g li{margin:4px 0}
@media(max-width:1100px){figure.g{grid-template-columns:1fr}}"""
shutil.rmtree("issues",ignore_errors=True);os.makedirs("issues/img")
lk=lambda us:" ".join(f'<a href="{E(u)}">{E(u)}</a>' for u in us if u.startswith("https://"))
tabs=[];allsrc={}
for q,lab,ks in TABS:
    figs=[]
    for k in ks.split():
        d=G[k];fn=f"{k}.webp";save_img(png[k],"issues/img/"+fn)
        b="".join(f'<li>{URL.sub(lambda m:f"<a href=\"{m[0]}\">{m[0]}</a>" if m[0].startswith("https://") else m[0],E(t))}</li>' for t,us in d["bul"])
        us=d["src"] or [u for _,x in d["bul"] for u in x]
        for u in us: allsrc.setdefault(u,1)
        if k!="01" and not d["bul"]: gaps.append(f"PNG {k}: no bullets")
        if not d["src"] and not any(x for _,x in d["bul"]): gaps.append(f"PNG {k}: no source")
        lab_=f'<p><i>{E(d["label"])}</i></p>' if d["label"] else ''
        sp=f'<p>Sources: {lk(d["src"])}</p>' if d["src"] else ''
        figs.append(f'<figure class="g"><img src="img/{fn}" alt="{E(d["title"])}" loading="lazy"><figcaption><span class="t">{E(d["title"])}</span><p>{E(d["caption"])}</p>{lab_}{"<ul>"+b+"</ul>" if b else ""}{sp}</figcaption></figure>')
    open(f"issues/T{q}.html","w").write(page(T,"T"+q,f'<h2 class="desk">{q}. {E(lab)}</h2>'+"\n".join(figs)));tabs.append(("T"+q,f"{q} {lab}"));print(q,lab,len(ks.split()))
for u in [u for _,d in G.items() for _,x in d["bul"] for u in x]: allsrc.setdefault(u,1)
for sl in list(P.slides)[39:]:
    for sh in sl.shapes:
        if sh.has_text_frame:
            for u in URL.findall(sh.text_frame.text):
                if u.endswith("..."): gaps.append("deck source truncated, omitted: "+u)
                else: allsrc.setdefault(u,1)
allsrc={u:1 for u in allsrc if u.startswith("https://")}
open("issues/sources.html","w").write(page(T,"sources",f'<h2 class="desk">Sources ({len(allsrc)} links)</h2><ul class="src">'+"".join(f'<li><a href="{E(u)}">{E(u)}</a></li>' for u in allsrc)+"</ul>"));tabs.append(("sources","Sources"))
index(T,tabs,"Sections",cur="renaiss")
print("\n".join(gaps))
